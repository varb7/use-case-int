from __future__ import annotations

import csv
import io
import os
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook
from pydantic import ValidationError
from streamlit.testing.v1 import AppTest

from verify import Q1FixtureProvider
from secure_settings import (API_KEY_ACCOUNT, MODEL_ACCOUNT, SERVICE,
                             remove_settings, resolve_api_key, resolve_model,
                             save_settings, settings_status, test_connection)
from workflow import (
    DraftProposal,
    EvidenceSection,
    EvidenceUse,
    GeminiProvider,
    JobStore,
    ProcessedAnswer,
    WorkflowEngine,
    export_csv,
    export_review_xlsx,
    load_evidence,
    parse_questionnaire,
    retrieve,
    route_owner,
    validate_draft,
)


ROOT = Path(__file__).parent
QUESTIONNAIRES = ROOT / "questionnaires"
TEST_DB = ROOT / "verification" / ".detailed-test.sqlite3"
UI_DB = ROOT / "verification" / ".ui-test.sqlite3"


def source(ref_id: str = "KB:test.md#control", *, source_date: date = date(2026, 8, 1)) -> EvidenceSection:
    return EvidenceSection(
        ref_id=ref_id,
        citation="test.md — Control",
        text="We encrypt customer data with AES-256.",
        owner="Security",
        source_date=source_date,
    )


def valid_draft(item: EvidenceSection | None = None) -> DraftProposal:
    item = item or source()
    return DraftProposal(
        status="Yes",
        answer="We encrypt customer data with AES-256.",
        evidence_uses=[EvidenceUse(ref_id=item.ref_id, quote=item.text)],
        owner=None,
        confidence="High",
        evidence_strength_reason="Direct source.",
    )


class StaticProvider:
    def __init__(self, draft: DraftProposal | None = None, error: Exception | None = None):
        self.result = draft or valid_draft()
        self.error = error
        self.calls = 0

    def draft(self, question, evidence):
        self.calls += 1
        if self.error:
            raise self.error
        return self.result


class ParsingTests(unittest.TestCase):
    def test_all_supplied_layouts_and_counts(self):
        actual = []
        for path in sorted(QUESTIONNAIRES.glob("*.csv")):
            parsed = parse_questionnaire(path.read_bytes(), path.name)
            actual.append((len(parsed.rows), parsed.question_column))
        self.assertEqual(actual, [(25, "question"), (30, "requirement"), (20, "requirement_text")])

    def test_original_columns_order_and_row_order_are_preserved(self):
        parsed = parse_questionnaire((QUESTIONNAIRES / "questionnaire-1-standard.csv").read_bytes(), "q1.csv")
        self.assertEqual(parsed.headers, ["ref", "section", "question", "vendor_response", "vendor_comments"])
        self.assertEqual(parsed.rows[0]["ref"], "Q1")
        self.assertEqual(parsed.rows[-1]["ref"], "Q25")

    def test_filename_is_reduced_to_basename(self):
        parsed = parse_questionnaire(b"question\nTest\n", "../../unsafe.csv")
        self.assertEqual(parsed.filename, "unsafe.csv")

    def test_explicit_custom_question_column_is_supported(self):
        parsed = parse_questionnaire(b"id,prompt\n1,Test question\n", "custom.csv", "prompt")
        self.assertEqual(parsed.question_column, "prompt")

    def test_rejects_oversized_file(self):
        with self.assertRaisesRegex(ValueError, "5 MB"):
            parse_questionnaire(b"x" * 5_000_001, "large.csv")

    def test_rejects_non_utf8(self):
        with self.assertRaisesRegex(ValueError, "UTF-8"):
            parse_questionnaire(b"question\n\xff\n", "bad.csv")

    def test_rejects_empty_file(self):
        with self.assertRaisesRegex(ValueError, "header"):
            parse_questionnaire(b"", "empty.csv")

    def test_rejects_duplicate_headers(self):
        with self.assertRaisesRegex(ValueError, "unique"):
            parse_questionnaire(b"question,question\none,two\n", "duplicate.csv")

    def test_rejects_missing_question_mapping(self):
        with self.assertRaisesRegex(ValueError, "Select"):
            parse_questionnaire(b"id,text\n1,Hello\n", "unknown.csv")

    def test_rejects_blank_question(self):
        with self.assertRaisesRegex(ValueError, "Every row"):
            parse_questionnaire(b"question,other\n,value\n", "blank.csv")

    def test_rejects_extra_csv_fields(self):
        with self.assertRaisesRegex(ValueError, "more fields"):
            parse_questionnaire(b"question\nvalue,extra\n", "extra.csv")

    def test_rejects_malformed_csv_quotes(self):
        with self.assertRaisesRegex(ValueError, "Malformed"):
            parse_questionnaire(b'question\n"unterminated\n', "quotes.csv")


class EvidenceAndRoutingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.evidence = load_evidence(ROOT)

    def test_loads_all_knowledge_and_approved_answers(self):
        self.assertEqual(len(self.evidence), 109)
        self.assertEqual(sum(item.ref_id.startswith("AA-") for item in self.evidence), 42)
        self.assertTrue(any("01-information-security-policy" in item.ref_id for item in self.evidence))

    def test_retrieval_is_bounded_and_finds_encryption(self):
        results = retrieve("Is customer data encrypted at rest?", self.evidence)
        self.assertLessEqual(len(results), 8)
        self.assertTrue(any("02-encryption-and-key-management" in item.ref_id for item in results))

    def test_no_keyword_overlap_returns_no_evidence(self):
        self.assertEqual(retrieve("zzzxxyyqqq", self.evidence), [])

    def test_lexical_morphology_miss_bypasses_model(self):
        provider = StaticProvider()
        result = WorkflowEngine([source()], provider).process("Describe encryption")
        self.assertEqual((result.status, provider.calls), ("Cannot answer", 0))
        self.assertIn("No relevant", result.evidence_strength_reason)

    def test_owner_routes_cover_all_teams(self):
        self.assertEqual(route_owner("Will you accept unlimited liability?"), "Legal")
        self.assertEqual(route_owner("Can the product be deployed on premise?"), "Product")
        self.assertEqual(route_owner("What discount is available?"), "Sales")
        self.assertEqual(route_owner("Describe encryption controls"), "Security")

    def test_retention_conflict_bypasses_model_and_routes_legal(self):
        provider = StaticProvider()
        result = WorkflowEngine(self.evidence, provider).process("For how long are consent records retained after collection?")
        self.assertEqual((result.status, result.owner, provider.calls), ("Cannot answer", "Legal", 0))
        self.assertIn("conflict", result.evidence_strength_reason.lower())

    def test_no_evidence_bypasses_model(self):
        provider = StaticProvider()
        result = WorkflowEngine(self.evidence, provider).process("zzzxxyyqqq")
        self.assertEqual((result.status, provider.calls), ("Cannot answer", 0))
        self.assertIn("No relevant", result.evidence_strength_reason)


class ValidationTests(unittest.TestCase):
    def setUp(self):
        self.item = source()
        self.by_id = {self.item.ref_id: self.item}

    def test_valid_supported_answer_passes(self):
        result = validate_draft("Is data encrypted?", valid_draft(self.item), self.by_id)
        self.assertEqual((result.status, result.validation_errors), ("Yes", []))

    def test_unknown_evidence_reference_is_blocked(self):
        draft = valid_draft(self.item).model_copy(update={"evidence_uses": [EvidenceUse(ref_id="missing", quote=self.item.text)]})
        result = validate_draft("Is data encrypted?", draft, self.by_id)
        self.assertEqual(result.status, "Cannot answer")
        self.assertIn("Unknown evidence", result.validation_errors[0])

    def test_non_exact_quote_is_blocked(self):
        draft = valid_draft(self.item).model_copy(update={"evidence_uses": [EvidenceUse(ref_id=self.item.ref_id, quote="AES-128") ]})
        result = validate_draft("Is data encrypted?", draft, self.by_id)
        self.assertEqual(result.status, "Cannot answer")
        self.assertTrue(any("exact source excerpt" in error for error in result.validation_errors))

    def test_substantive_answer_without_evidence_is_blocked(self):
        draft = valid_draft(self.item).model_copy(update={"evidence_uses": []})
        result = validate_draft("Is data encrypted?", draft, self.by_id)
        self.assertIn("Substantive statuses require evidence", result.validation_errors)

    def test_cannot_answer_without_owner_is_blocked(self):
        draft = DraftProposal(status="Cannot answer", answer="We cannot confirm this.", owner=None, confidence="Low")
        result = validate_draft("Can you confirm this?", draft, self.by_id)
        self.assertIn("Cannot answer requires an owner", result.validation_errors)

    def test_non_first_person_answer_is_blocked(self):
        draft = valid_draft(self.item).model_copy(update={"answer": "Customer data is encrypted."})
        result = validate_draft("Is data encrypted?", draft, self.by_id)
        self.assertTrue(any("first-person" in error for error in result.validation_errors))

    def test_more_than_four_sentences_is_blocked(self):
        draft = valid_draft(self.item).model_copy(update={"answer": "We do one. We do two. We do three. We do four. We do five."})
        result = validate_draft("Is data encrypted?", draft, self.by_id)
        self.assertTrue(any("1 to 4" in error for error in result.validation_errors))

    def test_stale_source_downgrades_confidence_and_adds_note(self):
        stale = source(source_date=date(2024, 1, 1))
        result = validate_draft("Is data encrypted?", valid_draft(stale), {stale.ref_id: stale}, today=date(2026, 9, 14))
        self.assertEqual(result.confidence, "Medium")
        self.assertIn("older than 12 months", result.reviewer_note)

    def test_invalid_status_is_rejected_by_schema(self):
        with self.assertRaises(ValidationError):
            DraftProposal(status="Maybe", answer="We are unsure.", confidence="Low")

    def test_gemini_schema_contains_no_empty_enum_value(self):
        def enums(value):
            if isinstance(value, dict):
                if "enum" in value:
                    yield value["enum"]
                for child in value.values():
                    yield from enums(child)
            elif isinstance(value, list):
                for child in value:
                    yield from enums(child)

        self.assertFalse(any("" in enum for enum in enums(DraftProposal.model_json_schema())))


class ProviderAndEngineTests(unittest.TestCase):
    def test_generation_requests_are_paced_and_429_retry_is_bounded(self):
        from google.genai import types
        from google.genai.errors import ClientError

        provider = object.__new__(GeminiProvider)
        provider.model, provider.types = "test-gemini", types
        provider._minimum_generation_interval = 4.1
        sample = valid_draft()
        limited = ClientError(429, {"error": {"code": 429, "message": "quota", "details": [
            {"@type": "type.googleapis.com/google.rpc.RetryInfo", "retryDelay": "7s"}]}})
        calls = SimpleNamespace(generate_content=unittest.mock.Mock(
            side_effect=[limited, SimpleNamespace(text=sample.model_dump_json()),
                         SimpleNamespace(text=sample.model_dump_json())]))
        provider.client = SimpleNamespace(models=calls)
        GeminiProvider._next_generation_at = 0
        with patch("workflow.time.monotonic", return_value=100.0), patch("workflow.time.sleep") as sleep:
            self.assertEqual(provider.draft("first", [source()]), sample)
            self.assertEqual(provider.draft("second", [source()]), sample)
        self.assertEqual(calls.generate_content.call_count, 3)
        sleep.assert_any_call(8.0)
        self.assertAlmostEqual(sum(call.args[0] for call in sleep.call_args_list), 16.2)
        GeminiProvider._next_generation_at = 0

    def test_generation_interval_configuration_is_validated(self):
        environment = {"GEMINI_API_KEY": "test", "GEMINI_MODEL": "test", "GEMINI_GENERATION_INTERVAL_SECONDS": "bad"}
        with patch.dict(os.environ, environment), self.assertRaisesRegex(RuntimeError, "must be a number"):
            GeminiProvider()

    def test_missing_google_key_is_reported_before_client_creation(self):
        with patch.dict(os.environ, {"GEMINI_API_KEY": "", "GOOGLE_API_KEY": ""}), patch("workflow.resolve_api_key", return_value=None), self.assertRaisesRegex(RuntimeError, "Google API key"):
            GeminiProvider()

    def test_google_api_key_alias_is_accepted_without_printing_key(self):
        with patch.dict(os.environ, {"GEMINI_API_KEY": "", "GOOGLE_API_KEY": "test-secret", "GEMINI_MODEL": "test-generation-model"}), patch("google.genai.Client") as client:
            provider = GeminiProvider()
        self.assertEqual(provider.model, "test-generation-model")
        self.assertEqual(client.call_args.kwargs["api_key"], "test-secret")

    def test_gemini_structured_response_contract(self):
        sample = valid_draft()

        class Models:
            def generate_content(self, **kwargs):
                self.kwargs = kwargs
                return SimpleNamespace(text=sample.model_dump_json())

        provider = object.__new__(GeminiProvider)
        from google.genai import types
        provider.model, provider.types = "test-gemini", types
        models = Models()
        provider.client = SimpleNamespace(models=models)
        self.assertEqual(provider.draft("test", [source()]), sample)
        self.assertEqual(models.kwargs["config"].response_mime_type, "application/json")
        self.assertIn("untrusted data", models.kwargs["config"].system_instruction)

    def test_empty_model_response_becomes_processing_failure(self):
        provider = object.__new__(GeminiProvider)
        from google.genai import types
        provider.model, provider.types = "test-gemini", types
        provider.client = SimpleNamespace(models=SimpleNamespace(generate_content=lambda **kwargs: SimpleNamespace(text="")))
        result = WorkflowEngine([source()], provider).process("Is customer data encrypted?")
        self.assertIn("Model returned no structured draft", result.processing_error)

    def test_invalid_model_json_becomes_processing_failure(self):
        provider = object.__new__(GeminiProvider)
        from google.genai import types
        provider.model, provider.types = "test-gemini", types
        provider.client = SimpleNamespace(models=SimpleNamespace(generate_content=lambda **kwargs: SimpleNamespace(text="not json")))
        result = WorkflowEngine([source()], provider).process("Is customer data encrypted?")
        self.assertIn("ValidationError", result.processing_error)

    def test_provider_exception_is_not_misreported_as_business_answer(self):
        result = WorkflowEngine([source()], StaticProvider(error=TimeoutError("timeout"))).process("Is customer data encrypted?")
        self.assertEqual(result.status, "Cannot answer")
        self.assertIn("TimeoutError", result.processing_error)
        self.assertIn("Processing failure", result.evidence_strength_reason)


class SecureSettingsTests(unittest.TestCase):
    def setUp(self):
        class DeleteError(Exception):
            pass

        self.values = {}
        self.backend = SimpleNamespace(
            errors=SimpleNamespace(PasswordDeleteError=DeleteError),
            get_password=lambda service, account: self.values.get((service, account)),
            set_password=lambda service, account, value: self.values.__setitem__((service, account), value),
            delete_password=self._delete,
        )

    def _delete(self, service, account):
        if (service, account) not in self.values:
            raise self.backend.errors.PasswordDeleteError()
        del self.values[(service, account)]

    def test_save_load_remove_and_never_store_blank(self):
        with patch("secure_settings._keyring", return_value=self.backend), patch.dict(
                os.environ, {"GEMINI_API_KEY": "", "GOOGLE_API_KEY": "", "GEMINI_MODEL": ""}):
            save_settings("  secret  ", "  test-model  ")
            self.assertEqual(resolve_api_key(), "secret")
            self.assertEqual(resolve_model(), "test-model")
            self.assertEqual(settings_status(), (True, True, "Windows Credential Manager"))
            save_settings(None, None)
            remove_settings()
            self.assertEqual(settings_status()[:2], (False, False))

    def test_environment_variables_take_precedence(self):
        self.values[(SERVICE, API_KEY_ACCOUNT)] = "saved-key"
        self.values[(SERVICE, MODEL_ACCOUNT)] = "saved-model"
        with patch("secure_settings._keyring", return_value=self.backend), patch.dict(
                os.environ, {"GEMINI_API_KEY": "environment-key", "GOOGLE_API_KEY": "", "GEMINI_MODEL": "environment-model"}):
            self.assertEqual(resolve_api_key(), "environment-key")
            self.assertEqual(resolve_model(), "environment-model")
            self.assertEqual(settings_status(), (True, True, "environment"))

    def test_connection_uses_secret_without_returning_it(self):
        self.values[(SERVICE, API_KEY_ACCOUNT)] = "never-display-this"
        self.values[(SERVICE, MODEL_ACCOUNT)] = "test-model"
        client = unittest.mock.Mock()
        client.return_value.models.get.return_value = SimpleNamespace(name="models/test-model")
        with patch("secure_settings._keyring", return_value=self.backend), patch.dict(
                os.environ, {"GEMINI_API_KEY": "", "GOOGLE_API_KEY": "", "GEMINI_MODEL": ""}), patch(
                    "google.genai.Client", client):
            self.assertEqual(test_connection(), "models/test-model")
        self.assertEqual(client.call_args.kwargs["api_key"], "never-display-this")


class StoreReviewAndExportTests(unittest.TestCase):
    def setUp(self):
        TEST_DB.parent.mkdir(exist_ok=True)
        if TEST_DB.exists():
            TEST_DB.unlink()
        self.store = JobStore(TEST_DB)

    def tearDown(self):
        if TEST_DB.exists():
            TEST_DB.unlink()

    def test_full_q1_job_persists_all_rows(self):
        questionnaire = parse_questionnaire((QUESTIONNAIRES / "questionnaire-1-standard.csv").read_bytes(), "q1.csv")
        evidence = load_evidence(ROOT)
        job_id = self.store.create_job(questionnaire)
        progress = []
        self.store.process_job(job_id, WorkflowEngine(evidence, Q1FixtureProvider()), progress=lambda done, total: progress.append((done, total)))
        self.assertEqual((len(self.store.rows(job_id)), progress[-1], self.store.job(job_id)["state"]), (25, (25, 25), "complete"))

    def test_failed_row_retries_without_regenerating_successful_row(self):
        questionnaire = parse_questionnaire(b"question\nIs customer data encrypted one?\nIs customer data encrypted two?\n", "retry.csv")
        item = source()

        class FirstPass:
            def __init__(self): self.calls = 0
            def draft(self, question, evidence):
                self.calls += 1
                if "two" in question: raise TimeoutError("timeout")
                return valid_draft(item)

        first = FirstPass()
        job_id = self.store.create_job(questionnaire)
        self.store.process_job(job_id, WorkflowEngine([item], first))
        self.assertEqual(self.store.job(job_id)["state"], "partial")
        retry = StaticProvider(valid_draft(item))
        self.store.process_job(job_id, WorkflowEngine([item], retry), failed_only=True)
        self.assertEqual((retry.calls, self.store.job(job_id)["state"]), (1, "complete"))

    def test_accept_then_edit_resets_review_state(self):
        questionnaire = parse_questionnaire(b"question\nEncryption\n", "review.csv")
        job_id = self.store.create_job(questionnaire)
        answer = ProcessedAnswer(**valid_draft().model_dump())
        self.store.save_review(job_id, 0, answer, accept=True)
        self.assertEqual(self.store.rows(job_id)[0]["review_state"], "Accepted")
        answer.answer = "We edited this answer."
        self.store.save_review(job_id, 0, answer, accept=False)
        self.assertEqual(self.store.rows(job_id)[0]["review_state"], "Needs review")

    def test_invalid_or_failed_answer_cannot_be_accepted(self):
        questionnaire = parse_questionnaire(b"question\nEncryption\n", "review.csv")
        job_id = self.store.create_job(questionnaire)
        answer = ProcessedAnswer(**valid_draft().model_dump(), validation_errors=["bad citation"])
        self.store.save_review(job_id, 0, answer, accept=True)
        self.assertEqual(self.store.rows(job_id)[0]["review_state"], "Needs review")

    def test_csv_export_preserves_partial_row_and_escapes_formulas(self):
        questionnaire = parse_questionnaire(b"question,comments\nUnknown,=2+2\n", "unsafe.csv")
        job_id = self.store.create_job(questionnaire)
        rows = list(csv.DictReader(io.StringIO(export_csv(self.store, job_id, []).decode("utf-8-sig"))))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["comments"], "'=2+2")
        self.assertEqual(rows[0]["status"], "")

    def test_cannot_answer_export_has_owner_and_blank_evidence(self):
        questionnaire = parse_questionnaire(b"question\nRetention?\n", "cannot.csv")
        job_id = self.store.create_job(questionnaire)
        answer = ProcessedAnswer(
            status="Cannot answer", answer="We cannot confirm this.", owner="Legal", confidence="Low",
            evidence_uses=[EvidenceUse(ref_id=source().ref_id, quote=source().text)],
        )
        self.store.save_review(job_id, 0, answer)
        row = next(csv.DictReader(io.StringIO(export_csv(self.store, job_id, [source()]).decode("utf-8-sig"))))
        self.assertEqual((row["owner"], row["evidence"]), ("Legal", ""))

    def test_review_workbook_contains_internal_fields_and_all_rows(self):
        questionnaire = parse_questionnaire(b"question\nEncryption\nUnprocessed\n", "review.csv")
        job_id = self.store.create_job(questionnaire)
        self.store.save_review(job_id, 0, ProcessedAnswer(**valid_draft().model_dump()))
        workbook = load_workbook(io.BytesIO(export_review_xlsx(self.store, job_id, [source()])), read_only=True)
        sheet = workbook["Review"]
        headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        self.assertEqual(sheet.max_row, 3)
        self.assertIn("processing_error", headers)
        self.assertIn("evidence_excerpts", headers)


class StreamlitShellTests(unittest.TestCase):
    def test_failure_and_routing_filters_preserve_row_identity_and_exports(self):
        questionnaire = parse_questionnaire(b"question\nBroken request\nLegal terms\nSecurity review\nSupported\n", "filters.csv")
        job = self.store.create_job(questionnaire)
        self.store.save_review(job, 0, ProcessedAnswer(status="Cannot answer", answer="We could not process this.",
                              owner="Security", confidence="Low", processing_error="429 quota exceeded"))
        for index, owner in [(1, "Legal"), (2, "Security")]:
            self.store.save_review(job, index, ProcessedAnswer(status="Cannot answer", answer="We need owner review.",
                                  owner=owner, confidence="Low", escalation_reason=f"Requires {owner} approval"))
        self.store.save_review(job, 3, ProcessedAnswer(**valid_draft().model_dump()))
        with patch("workflow.JobStore", return_value=self.store):
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            select = lambda label: next(box for box in app.selectbox if box.label == label)
            select("Show questions").select("Failed").run(timeout=20)
            self.assertEqual(len(select("Question").options), 1)
            self.assertTrue(select("Question").value.startswith("1:"))
            self.assertTrue(any("429 quota exceeded" in error.value for error in app.error))
            select("Show questions").select("Owner-routed").run(timeout=20)
            self.assertEqual(len(select("Question").options), 2)
            select("Filter by owner").select("Legal").run(timeout=20)
            self.assertTrue(select("Question").value.startswith("2:"))
            self.assertEqual(select("Owner").value, "Legal")
            self.assertTrue(any("Requires Legal approval" in warning.value for warning in app.warning))
            self.assertEqual(len(app.get("download_button")), 2)
            select("Show questions").select("All questions").run(timeout=20)
            select("Filter by owner").select("All owners").run(timeout=20)
            select("Question").select(next(v for v in select("Question").options if v.startswith("4:"))).run(timeout=20)
            self.assertFalse(any(area.label in ("Exact source excerpt", "Supported answer text") for area in app.text_area))
            self.assertTrue(any("Evidence and checks" in item.value for item in app.markdown))
            self.assertFalse(app.exception)
        self.assertEqual(len(list(csv.DictReader(io.StringIO(export_csv(self.store, job, []).decode("utf-8-sig"))))), 4)

    def setUp(self):
        if UI_DB.exists():
            UI_DB.unlink()
        self.store = JobStore(UI_DB)
        self.credentials = patch("secure_settings._read", return_value=None)
        self.credentials.start()

    def tearDown(self):
        self.credentials.stop()
        if UI_DB.exists():
            UI_DB.unlink()

    def test_app_starts_with_upload_control(self):
        with patch.dict(os.environ, {"GEMINI_API_KEY": "", "GOOGLE_API_KEY": ""}), patch("workflow.JobStore", return_value=self.store):
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
        self.assertFalse(app.exception)
        self.assertEqual(app.title[0].value, "Guided questionnaire review")
        self.assertEqual(len(app.get("file_uploader")), 1)
        self.assertTrue(any(button.label == "Save settings" for button in app.button))
        self.assertTrue(any(field.label == "Google API key" for field in app.text_input))
        self.assertIn('type="password"', (ROOT / "app.py").read_text(encoding="utf-8"))

    def test_valid_q1_upload_is_previewed_and_key_gate_is_visible(self):
        data = (QUESTIONNAIRES / "questionnaire-1-standard.csv").read_bytes()
        with patch.dict(os.environ, {"GEMINI_API_KEY": "", "GOOGLE_API_KEY": ""}), patch("workflow.JobStore", return_value=self.store):
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            app.get("file_uploader")[0].upload("q1.csv", data, "text/csv").run(timeout=20)
            generate = next(button for button in app.button if button.label == "Generate draft")
            generate.click().run(timeout=20)
        self.assertFalse(app.exception)
        self.assertTrue(any("Complete Google settings" in error.value for error in app.error))

    def test_invalid_upload_shows_useful_error(self):
        with patch.dict(os.environ, {"GEMINI_API_KEY": "", "GOOGLE_API_KEY": ""}), patch("workflow.JobStore", return_value=self.store):
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            app.get("file_uploader")[0].upload("bad.csv", b"id,text\n1,hello\n", "text/csv").run(timeout=20)
        self.assertFalse(app.exception)
        self.assertTrue(any("Select the column" in error.value for error in app.error))

    def test_unrecognized_layout_is_rejected_before_manual_mapping(self):
        with patch.dict(os.environ, {"GEMINI_API_KEY": "", "GOOGLE_API_KEY": ""}), patch("workflow.JobStore", return_value=self.store):
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            app.get("file_uploader")[0].upload("custom.csv", b"id,prompt\n1,Question text\n", "text/csv").run(timeout=20)
        self.assertTrue(any("Select the column" in error.value for error in app.error))
        self.assertFalse(any(box.label == "Question column" for box in app.selectbox))

    def test_mocked_gemini_q1_runs_through_review_and_export_controls(self):
        data = (QUESTIONNAIRES / "questionnaire-1-standard.csv").read_bytes()
        evidence = load_evidence(ROOT)
        index = SimpleNamespace(
            sync=lambda sections: None,
            search=lambda question, limit=8: [(s, 1.0) for s in retrieve(question, evidence, limit)],
        )
        with (
            patch.dict(os.environ, {"GEMINI_API_KEY": "test-only", "GOOGLE_API_KEY": "", "GEMINI_MODEL": "test-model"}),
            patch("workflow.JobStore", return_value=self.store),
            patch("semantic.GeminiProvider", return_value=Q1FixtureProvider()),
            patch("semantic.GoogleEmbeddings"),
            patch("semantic.SemanticIndex", return_value=index),
        ):
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            app.get("file_uploader")[0].upload("q1.csv", data, "text/csv").run(timeout=20)
            next(button for button in app.button if button.label == "Generate draft").click().run(timeout=30)
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        self.assertEqual(len(self.store.rows(self.store.list_jobs()[0]["id"])), 25)
        self.assertEqual(len(app.get("download_button")), 2)


if __name__ == "__main__":
    output = io.StringIO()
    result = unittest.TextTestRunner(stream=output, verbosity=2).run(
        unittest.defaultTestLoader.loadTestsFromModule(__import__(__name__))
    )
    text = output.getvalue()
    print(text, end="")
    (ROOT / "verification").mkdir(exist_ok=True)
    (ROOT / "verification" / "DETAILED_TEST_OUTPUT.txt").write_text(text, encoding="utf-8")
    raise SystemExit(0 if result.wasSuccessful() else 1)
