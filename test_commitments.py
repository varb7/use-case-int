"""Offline commitment-policy tests. Run directly to save a reproducible report."""
import csv
import io
import json
import tempfile
import time
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import Mock

from openpyxl import load_workbook

from commitments import AUTHORITY, commitment_request
from workflow import (DraftProposal, EvidenceSection, EvidenceUse, JobStore,
                      ProcessedAnswer, WorkflowEngine, commitment_guard,
                      export_csv, export_review_xlsx, guarded_rows, load_evidence,
                      parse_questionnaire, validate_draft)

ROOT = Path(__file__).parent
TODAY = date(2026, 9, 15)
LIABILITY = "KB:09-privacy-and-dpa-terms.md#liability"


def proposal(answer="We guarantee a 99.99% uptime SLA.", status="Yes", uses=None):
    return DraftProposal(status=status, answer=answer, confidence="High",
                         owner="Security", evidence_uses=uses or [])


class CommitmentTests(unittest.TestCase):
    def setUp(self):
        self.evidence = load_evidence(ROOT)
        self.sources = {s.ref_id: s for s in self.evidence}

    def test_new_commitment_questions(self):
        questions = [
            "Will you accept unlimited liability for data protection breaches?",
            "Do you commit to a 99.99% monthly uptime SLA with service credits?",
            "Are you HIPAA compliant and will you sign a BAA?",
            "Do you agree to our standard DPA in place of yours?",
            "Can you guarantee deletion within seven days?",
            "Will you indemnify us for a breach?",
            "Please agree to our contractual terms.",
            "You must accept our liability terms.",
            "Would you be willing to sign our BAA?",
            "WILL  YOU\nSIGN our BAA?", "Ｗｉｌｌ you sign our BAA?",
        ]
        for question in questions:
            with self.subTest(question=question):
                self.assertEqual(commitment_request(question), ("C1-new-commitment", "Legal"))

    def test_factual_questions_not_preblocked(self):
        for question in [
            "What is your standard liability cap?",
            "What availability target is in your standard terms?",
            "What is your current deletion policy?",
            "Within what timeframe will you notify us of a personal data breach?",
            "What breach-notification period does your DPA state?",
            "Please provide your latest SOC 2 report.",
            "Can we export our data?", "Do you support SSO?",
            "Describe your availability commitments and historical performance.",
            "Confirm your position on liability caps and indemnities.",
        ]:
            with self.subTest(question=question):
                self.assertIsNone(commitment_request(question))

    def test_sales_and_legal_precedence(self):
        self.assertEqual(commitment_request("Will you guarantee a fixed price?")[1], "Sales")
        self.assertEqual(commitment_request("Will you agree to a discount and unlimited liability?")[1], "Legal")

    def test_no_retrieval_or_provider_call_for_commitment(self):
        provider, retriever = Mock(), Mock()
        result = WorkflowEngine(self.evidence, provider, retriever).process("Will you sign a BAA?")
        provider.draft.assert_not_called()
        retriever.assert_not_called()
        self.assertEqual((result.status, result.owner, result.evidence_uses), ("Cannot answer", "Legal", []))
        self.assertFalse(result.processing_error)

    def test_mixed_question_routes_whole_row(self):
        result = commitment_guard("Do you encrypt data and will you guarantee zero breaches?")
        self.assertEqual(result.status, "Cannot answer")

    def test_all_supplied_questionnaire_rules(self):
        expected = [[], [3, 6, 20, 22, 23], []]
        for path, rows in zip(sorted((ROOT / "questionnaires").glob("*.csv")), expected):
            q = parse_questionnaire(path.read_bytes(), path.name)
            actual = [i + 1 for i, row in enumerate(q.rows) if commitment_request(row[q.question_column])]
            self.assertEqual(actual, rows, path.name)

    def test_q1_fixture_still_processes_all_rows(self):
        from verify import Q1FixtureProvider
        q = parse_questionnaire((ROOT / "questionnaires/questionnaire-1-standard.csv").read_bytes(), "q1.csv")
        engine = WorkflowEngine(self.evidence, Q1FixtureProvider(), lambda q, evidence: evidence)
        results = [engine.process(row[q.question_column]) for row in q.rows]
        self.assertEqual(len(results), 25)
        self.assertFalse(any(r.processing_error or r.validation_errors for r in results))
        self.assertFalse(any("C1-" in r.reviewer_note or "C3-" in r.reviewer_note for r in results))

    def test_authority_is_application_owned_and_serialized(self):
        source = self.sources[LIABILITY]
        self.assertEqual(source.model_dump()["authority"], "existing_standard_term")
        forged = EvidenceSection(ref_id="unregistered", citation="fake", text=source.text,
                                 authority="existing_standard_term")
        self.assertEqual(forged.authority, "factual_policy")

    def test_changed_revision_loses_authority(self):
        changed = self.sources[LIABILITY].model_copy(update={"text": "We accept unlimited liability."})
        self.assertEqual(changed.authority, "factual_policy")
        self.assertEqual(self.sources[LIABILITY].authority, "existing_standard_term")

    def test_negotiation_sections_do_not_grant_permission(self):
        source = self.sources["KB:08-business-continuity.md#availability-commitment"]
        self.assertEqual(source.authority, "negotiation_only")
        draft = proposal(uses=[EvidenceUse(ref_id=source.ref_id, quote=source.text)])
        result = validate_draft("Describe availability.", draft, self.sources, TODAY)
        self.assertEqual(result.status, "Cannot answer")
        self.assertIn("C3-", result.validation_errors[0])

    def test_promises_blocked_regardless_of_status(self):
        for status in ("Yes", "No", "Partial", "Not applicable", "Cannot answer"):
            with self.subTest(status=status):
                result = validate_draft("Describe availability.", proposal(status=status), self.sources, TODAY)
                self.assertEqual(result.owner, "Legal")
                self.assertTrue(result.validation_errors)
                self.assertNotIn("99.99", result.answer)

    def test_affirmative_promise_variants(self):
        for answer in ["We'll sign your BAA.", "We shall indemnify you.",
                       "We accept unlimited liability.", "We are committed to your SLA.",
                       "We can commit to a 99.99% uptime SLA.",
                       "We sign your BAA.", "We will not lose any data.",
                       "We will not only provide service credits but also indemnify you.",
                       "Our SLA guarantees 99.99% uptime.", "We provide service credits.",
                       "We do not guarantee uptime, but we agree to service credits.",
                       "We encrypt data. We guarantee zero breaches."]:
            with self.subTest(answer=answer):
                self.assertIsNotNone(commitment_guard("Describe our policy.", proposal(answer), self.sources, TODAY))

    def test_factual_and_negative_statements_not_promise_blocked(self):
        for answer in ["We encrypt data with AES-256.", "We do not guarantee uninterrupted service.",
                       "We cannot agree to new terms without Legal approval.",
                       "We will not sign your BAA.",
                       "Our SLA is not a contractual guarantee.",
                       "Our availability target is 99.9 percent, not a contractual SLA."]:
            with self.subTest(answer=answer):
                self.assertIsNone(commitment_guard("Describe policy.", proposal(answer), self.sources, TODAY))

    def standard_draft(self, ref=LIABILITY):
        entry = AUTHORITY[ref]
        return proposal(entry[2], uses=[EvidenceUse(ref_id=ref, quote=entry[3], supported_claim=entry[2])])

    def test_exact_supported_standard_terms_allowed(self):
        for ref, entry in AUTHORITY.items():
            if entry[1] == "existing_standard_term":
                with self.subTest(ref=ref):
                    result = validate_draft("Describe existing standard terms.", self.standard_draft(ref), self.sources, TODAY)
                    self.assertEqual(result.status, "Yes")
                    self.assertFalse(result.validation_errors)

    def test_standard_source_cannot_authorize_new_request(self):
        result = validate_draft("Will you accept our liability terms?", self.standard_draft(), self.sources, TODAY)
        self.assertIn("C1-", result.reviewer_note)
        self.assertEqual(result.status, "Cannot answer")

    def test_changed_missing_stale_or_undated_standard_source_blocks(self):
        original = self.sources[LIABILITY]
        for change in ({"text": original.text + " Revised."}, {"source_date": date(2020, 1, 1)},
                       {"source_date": None}, {"source_date": date(2030, 1, 1)}):
            with self.subTest(change=change):
                sources = {LIABILITY: original.model_copy(update=change)}
                result = validate_draft("Describe existing terms.", self.standard_draft(), sources, TODAY)
                self.assertIn("C3-", result.reviewer_note)
        self.assertIsNotNone(commitment_guard("Describe terms.", self.standard_draft(), {}, TODAY))

    def test_wrong_excerpt_does_not_authorize_standard_term(self):
        draft = self.standard_draft()
        draft.evidence_uses[0].quote = "Changes to liability caps"
        result = validate_draft("Describe terms.", draft, self.sources, TODAY)
        self.assertIn("C3-", result.reviewer_note)

    def test_extended_or_modified_standard_term_is_blocked(self):
        for answer in [AUTHORITY[LIABILITY][2].replace("12 months", "24 months"),
                       AUTHORITY[LIABILITY][2] + " We also guarantee service credits.",
                       AUTHORITY[LIABILITY][2][:-1] + " and we waive that cap."]:
            with self.subTest(answer=answer):
                draft = self.standard_draft()
                draft.answer = answer
                self.assertIsNotNone(commitment_guard("Describe existing terms.", draft, self.sources, TODAY))

    def test_valid_authority_still_requires_retrieved_reference(self):
        result = validate_draft("Describe existing terms.", self.standard_draft(), self.sources, TODAY, allowed_refs=[])
        self.assertTrue(result.validation_errors)

    def test_no_llm_judge_on_postdraft_failure(self):
        provider = Mock()
        provider.draft.return_value = proposal()
        result = WorkflowEngine(self.evidence, provider, lambda q, e: e).process("Describe availability.")
        self.assertEqual(provider.draft.call_count, 1)
        self.assertIn("C3-", result.reviewer_note)

    def test_sales_promise_routing(self):
        result = commitment_guard("Describe pricing.", proposal("We guarantee a fixed price."), self.sources, TODAY)
        self.assertEqual(result.owner, "Sales")

    def test_review_and_legacy_exports_recheck_commitments(self):
        with tempfile.TemporaryDirectory(prefix="commitment-tests-") as directory:
            store = JobStore(Path(directory) / "test.sqlite3")
            q = parse_questionnaire(b"id,question\n1,Will you sign our BAA?\n2,Describe availability.\n", "guard.csv")
            job = store.create_job(q)
            unsafe = ProcessedAnswer(**proposal().model_dump())
            store.save_review(job, 0, unsafe, accept=True, evidence=self.evidence)
            self.assertEqual(store.rows(job)[0]["result"].status, "Cannot answer")
            store.save_review(job, 1, unsafe, accept=True, evidence=self.evidence)
            self.assertEqual(store.rows(job)[1]["review_state"], "Needs review")
            # Simulate unsafe records saved before these guards existed.
            with store._connect() as db:
                db.execute("UPDATE answers SET result_json=?, review_state='Accepted' WHERE job_id=?",
                           (unsafe.model_dump_json(), job))
            exported = list(csv.DictReader(io.StringIO(export_csv(store, job, self.evidence).decode("utf-8-sig"))))
            self.assertEqual([r["status"] for r in exported], ["Cannot answer"] * 2)
            self.assertEqual([r["owner"] for r in exported], ["Legal"] * 2)
            self.assertTrue(all(not r["evidence"] for r in exported))
            book = load_workbook(io.BytesIO(export_review_xlsx(store, job, self.evidence)))
            headers = [c.value for c in book.active[1]]
            for row in book.active.iter_rows(min_row=2, values_only=True):
                self.assertEqual(row[headers.index("status")], "Cannot answer")
                self.assertNotIn("99.99", row[headers.index("answer")])
                self.assertEqual(row[headers.index("review_state")], "Needs review")
            displayed = guarded_rows(store, job, self.evidence)
            self.assertTrue(all(r["result"].owner == "Legal" for r in displayed))
            self.assertTrue(all(r["review_state"] == "Needs review" for r in displayed))
            # Export does not mutate the user's saved record.
            self.assertEqual(store.rows(job)[0]["result"].status, "Yes")

    def test_guarded_view_preserves_failures_and_accepted_routing(self):
        with tempfile.TemporaryDirectory(prefix="commitment-tests-") as directory:
            store = JobStore(Path(directory) / "test.sqlite3")
            q = parse_questionnaire(b"question\nWill you sign a BAA?\nWill you sign a BAA?\n", "guard.csv")
            job = store.create_job(q)
            routed = commitment_guard("Will you sign a BAA?")
            store.save_review(job, 0, routed, accept=True)
            failure = ProcessedAnswer(status="Cannot answer", answer="We could not process this row.",
                                      owner="Security", confidence="Low", processing_error="Timeout")
            with store._connect() as db:
                db.execute("UPDATE answers SET result_json=? WHERE job_id=? AND row_index=1",
                           (failure.model_dump_json(), job))
            rows = guarded_rows(store, job, self.evidence)
            self.assertEqual(rows[0]["review_state"], "Accepted")
            self.assertEqual(rows[1]["result"].processing_error, "Timeout")


if __name__ == "__main__":
    stream = io.StringIO()
    suite = unittest.defaultTestLoader.loadTestsFromNames(["test_commitments", "test_semantic", "test_detailed"])
    start = time.monotonic()
    result = unittest.TextTestRunner(stream=stream, verbosity=2).run(suite)
    report = {"run_date": date.today().isoformat(), "mode": "offline; no live Gemini calls",
              "tests_run": result.testsRun, "failures": len(result.failures), "errors": len(result.errors),
              "seconds": round(time.monotonic() - start, 3), "questionnaires": {}}
    for path in sorted((ROOT / "questionnaires").glob("*.csv")):
        q = parse_questionnaire(path.read_bytes(), path.name)
        report["questionnaires"][path.name] = [
            {"row": i + 1, "question": row[q.question_column], "pre_guard": commitment_request(row[q.question_column])}
            for i, row in enumerate(q.rows)]
    target = ROOT / "verification"
    target.mkdir(exist_ok=True)
    (target / "commitment-tests.txt").write_text(stream.getvalue(), encoding="utf-8")
    (target / "commitment-results.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(stream.getvalue())
    print(json.dumps({k: v for k, v in report.items() if k != "questionnaires"}, indent=2))
    raise SystemExit(0 if result.wasSuccessful() else 1)
