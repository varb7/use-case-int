"""One-command supplied-pack verification. Uses a deterministic test double, never a second production provider."""

from __future__ import annotations

import csv
import io
from collections import Counter
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from workflow import (
    DraftProposal,
    EvidenceSection,
    EvidenceUse,
    GeminiProvider,
    JobStore,
    Questionnaire,
    WorkflowEngine,
    export_csv,
    export_review_xlsx,
    load_evidence,
    parse_questionnaire,
)


ROOT = Path(__file__).parent


RULES = {
    "legal entity": ("Cannot answer", [], "We cannot provide our legal entity name and headquarters address from the supplied approved material.", "Legal"),
    "ISO 27001": ("Yes", ["07-certifications-audits.md#current"], "We are certified against ISO/IEC 27001:2022 under certificate IS 771204, valid until 30 November 2027.", ""),
    "SOC 2 Type II": ("Yes", ["07-certifications-audits.md#current"], "We have a SOC 2 Type II report covering Security and Availability for 1 May 2025 through 30 April 2026.", ""),
    "dedicated information security": ("Yes", ["01-information-security-policy.md#security-team"], "We have a dedicated four-person Security function led by our Head of Security.", ""),
    "encrypted at rest": ("Yes", ["02-encryption-and-key-management.md#data-at-rest"], "We encrypt all customer data at rest with AES-256.", ""),
    "encrypted in transit": ("Yes", ["02-encryption-and-key-management.md#data-in-transit"], "We use TLS 1.2 or higher for external traffic and prefer TLS 1.3.", ""),
    "keys managed": ("Yes", ["02-encryption-and-key-management.md#key-management"], "We manage encryption keys in Google Cloud KMS and rotate them automatically every 90 days.", ""),
    "bring our own encryption": ("No", ["02-encryption-and-key-management.md#key-management"], "We do not offer customer-managed encryption keys in the standard product.", ""),
    "most recent third-party penetration": ("Cannot answer", [], "We cannot confirm the date of our current penetration test from sufficiently current supplied evidence.", "Security"),
    "remediation SLAs": ("Partial", ["10-product-architecture.md#vulnerability-management"], "We use remediation targets of 7 days for Critical and 30 days for High vulnerabilities; these are targets rather than contractual SLAs.", ""),
    "multi-factor authentication": ("Yes", ["03-access-control.md#authentication"], "We enforce MFA for all staff and require hardware security keys for production access.", ""),
    "access rights to production reviewed": ("Yes", ["03-access-control.md#authorisation"], "We review production access quarterly.", ""),
    "access revoked": ("Yes", ["03-access-control.md#authorisation"], "We revoke access within 24 hours of termination through automated deprovisioning.", ""),
    "data hosted": ("Yes", ["06-data-residency-and-retention.md#hosting-regions"], "We host customer data in the EU by default, and Enterprise customers may select US hosting at account creation.", ""),
    "personal data of our website visitors": ("Yes", ["AA-021"], "We store pseudonymous consent details and process IP addresses transiently without storing them.", ""),
    "contract termination": ("Yes", ["06-data-residency-and-retention.md#deletion-on-termination"], "We delete customer data from primary storage within 30 days of termination and from backups within a further 30 days.", ""),
    "notify us of a personal data breach": ("Yes", ["04-incident-response.md#customer-notification"], "We notify affected customers without undue delay and within 48 hours of confirming a personal data breach.", ""),
    "reportable data breach": ("No", ["04-incident-response.md#history"], "We have not had a reportable personal data breach in the last 36 months.", ""),
    "24/7 security incident": ("Yes", ["04-incident-response.md#detection"], "We route security alerts to a 24/7 on-call Security Engineer.", ""),
    "RTO and RPO": ("Yes", ["08-business-continuity.md#recovery-objectives"], "We have a four-hour RTO and a one-hour RPO for the Admin Interface and APIs.", ""),
    "test your disaster recovery": ("Yes", ["08-business-continuity.md#dr-testing"], "We perform a full disaster recovery exercise annually; our November 2025 exercise met its targets.", ""),
    "Data Protection Officer": ("Yes", ["09-privacy-and-dpa-terms.md#data-protection-officer"], "We have appointed an external Data Protection Officer, reachable at dpo@.", ""),
    "List your subprocessors": ("Yes", ["05-subprocessors.md#current-subprocessors", "05-subprocessors.md#changes-to-the-list"], "We list our current subprocessors and notify customers at least 30 days before a new subprocessor begins processing.", ""),
    "used to train": ("No", ["AA-033"], "We do not use customer or end-user data to train machine-learning models.", ""),
}


class Q1FixtureProvider:
    """Deterministic test double for repeatable plumbing verification."""

    def draft(self, question: str, evidence: list[EvidenceSection]) -> DraftProposal:
        rule = next((value for key, value in RULES.items() if key.lower() in question.lower()), None)
        if not rule:
            raise AssertionError(f"No Q1 fixture rule for: {question}")
        status, ref_fragments, answer, owner = rule
        uses = []
        for fragment in ref_fragments:
            source = next((item for item in evidence if fragment in item.ref_id), None)
            if not source:
                raise AssertionError(f"Retrieval missed {fragment} for: {question}")
            uses.append(EvidenceUse(ref_id=source.ref_id, quote=source.text))
        return DraftProposal(
            status=status,
            answer=answer,
            evidence_uses=uses,
            owner=owner or None,
            confidence="High" if status not in ("Cannot answer", "Partial") else ("Low" if status == "Cannot answer" else "Medium"),
            reviewer_note="",
            evidence_strength_reason="Direct supplied source." if uses else "No sufficient supplied evidence.",
            what_needs_checking="Assigned owner must provide current evidence." if status == "Cannot answer" else "",
            escalation_reason="Material information is absent or not current." if status == "Cannot answer" else "",
        )


def load_questionnaires() -> list[Questionnaire]:
    return [parse_questionnaire(path.read_bytes(), path.name) for path in sorted((ROOT / "questionnaires").glob("*.csv"))]


def main() -> None:
    questionnaires = load_questionnaires()
    assert [(len(q.rows), q.question_column) for q in questionnaires] == [
        (25, "question"), (30, "requirement"), (20, "requirement_text")
    ]
    evidence = load_evidence(ROOT)
    from google.genai import types

    sample = DraftProposal(
        status="Cannot answer",
        answer="We cannot provide a supported answer.",
        owner="Security",
        confidence="Low",
    )

    class FakeModels:
        def generate_content(self, **kwargs):
            assert kwargs["config"].response_mime_type == "application/json"
            assert kwargs["config"].response_schema is DraftProposal
            return SimpleNamespace(text=sample.model_dump_json())

    gemini = object.__new__(GeminiProvider)
    gemini.model, gemini.types = "test-gemini", types
    gemini.client = SimpleNamespace(models=FakeModels())
    assert gemini.draft("test", evidence[:1]) == sample
    try:
        parse_questionnaire(b"question\nvalue,extra\n", "malformed.csv")
    except ValueError:
        pass
    else:
        raise AssertionError("Malformed CSV with extra fields was accepted")
    output_dir = ROOT / "verification"
    output_dir.mkdir(exist_ok=True)
    test_db = output_dir / ".verify.sqlite3"
    if test_db.exists():
        test_db.unlink()
    try:
        store = JobStore(test_db)
        job_id = store.create_job(questionnaires[0])
        store.process_job(job_id, WorkflowEngine(evidence, Q1FixtureProvider()))
        rows = store.rows(job_id)
        assert len(rows) == 25
        failures = [
            (row["row_index"] + 1, row["result"].processing_error, row["result"].validation_errors)
            for row in rows if row["result"].processing_error or row["result"].validation_errors
        ]
        assert not failures, failures
        counts = Counter(row["result"].status for row in rows)
        assert counts == {"Yes": 18, "No": 3, "Cannot answer": 3, "Partial": 1}
        assert rows[13]["result"].owner == "Legal" and "conflict" in rows[13]["result"].evidence_strength_reason.lower()

        first = rows[0]["result"]
        store.save_review(job_id, 0, first, accept=True)
        assert store.rows(job_id)[0]["review_state"] == "Accepted"
        first.answer += " We edited this answer."
        store.save_review(job_id, 0, first, accept=False)
        assert store.rows(job_id)[0]["review_state"] == "Needs review"

        csv_bytes = export_csv(store, job_id, evidence)
        exported = list(csv.DictReader(io.StringIO(csv_bytes.decode("utf-8-sig"))))
        assert len(exported) == 25 and list(exported[0])[-6:] == list(("status", "answer", "evidence", "owner", "confidence", "reviewer_note"))
        assert exported[13]["evidence"] == "" and exported[13]["owner"] == "Legal"
        xlsx_bytes = export_review_xlsx(store, job_id, evidence)
        workbook = load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
        assert workbook["Review"].max_row == 26

        unsafe = parse_questionnaire(b"question,comments\nUnknown question,=2+2\n", "unsafe.csv")
        unsafe_job = store.create_job(unsafe)
        unsafe_rows = list(csv.DictReader(io.StringIO(export_csv(store, unsafe_job, evidence).decode("utf-8-sig"))))
        assert len(unsafe_rows) == 1 and unsafe_rows[0]["comments"] == "'=2+2" and unsafe_rows[0]["status"] == ""

        (output_dir / "q1-fixture-output.csv").write_bytes(csv_bytes)
        (output_dir / "q1-fixture-review.xlsx").write_bytes(xlsx_bytes)
        report = f"""# Verification results

Run: `{datetime_now()}` with Python fixture provider (not a live Gemini call).

- Questionnaire parsing: Q1 25/25 rows, Q2 30/30 rows, Q3 20/20 rows; all question columns detected correctly.
- Gemini adapter contract: structured JSON/Pydantic request and response parsing passed without a network call.
- Complete Q1 plumbing: 25/25 processed, CSV 25/25 rows, review workbook 25/25 rows.
- Fixture statuses: Yes {counts['Yes']}, No {counts['No']}, Partial {counts['Partial']}, Cannot answer {counts['Cannot answer']}.
- Validation: 0 citation/quote/format errors and 0 processing failures.
- Conflict behavior: Q14 routed to Legal because the supplied sources state both 12 and 36 months.
- Review behavior: acceptance persisted; a subsequent edit reset the row to Needs review.
- Partial/safe export: an unprocessed row remained present and a formula-leading input was escaped as text.

Limitations: this command deliberately uses test doubles, regardless of environment keys. Live drafting quality, API latency, token cost and live-model status distribution were not measured. This command checks Q2/Q3 parsing/mapping only; other suites cover retrieval and commitment rules. Fixture outputs are not prospect-ready model results. See `VERIFICATION_REPORT.md` for consolidated current coverage.
"""
        (output_dir / "RESULTS.md").write_text(report, encoding="utf-8")
    finally:
        if test_db.exists():
            test_db.unlink()
    print(report)


def datetime_now() -> str:
    from datetime import datetime
    return datetime.now().isoformat(timespec="seconds")


if __name__ == "__main__":
    main()
